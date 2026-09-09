"""クライアント記入用テンプレートが、そのまま取り込めること。"""
from __future__ import annotations

import io
from datetime import date, timedelta

from openpyxl import load_workbook

from app.importer.template import build_template, template_bytes

from .conftest import TODAY


def test_template_has_a_legend_and_one_example_row():
    book = build_template(TODAY)
    sheet = book["課題リスト"]

    assert [c.value for c in sheet[5]][:4] == [
        "No", "課題・気になっていること（必須）", "関連するタスク・工程", "重要度",
    ]
    # 記入例は1行だけ、グレーで見分けが付くこと
    assert sheet["A6"].value == "例"
    assert sheet["A6"].font.italic is True
    assert sheet["B7"].value is None  # 記入欄は空

    guide = book["記入のしかた"]
    text = "\n".join(str(row[0].value or "") for row in guide.iter_rows(max_col=1))
    assert "必須は「課題・気になっていること」の列だけです" in text
    assert "自動で除外します" in text


def test_template_offers_a_severity_dropdown():
    sheet = build_template(TODAY)["課題リスト"]
    lists = [dv for dv in sheet.data_validations.dataValidation if dv.type == "list"]
    assert lists and lists[0].formula1 == '"高,中,低"'


def test_template_is_recognised_as_an_issue_list(client):
    """記入例だけが入った状態でも、課題リストとして判定される。"""
    analyzed = client.post(
        "/api/imports/analyze",
        files={"file": ("template.xlsx", template_bytes(TODAY), "application/octet-stream")},
    ).json()

    assert analyzed["header_row"] == 4  # タイトル3行の下
    assert analyzed["mapping"]["issue"] == "課題・気になっていること(必須)"  # 全角括弧はNFKCで半角化
    assert analyzed["mapping"]["priority"] == "重要度"
    assert analyzed["content_kind"] == "issues"


def test_a_filled_in_template_imports_end_to_end(client, project, factory):
    factory.task("システム要件定義")

    book = build_template(TODAY)
    sheet = book["課題リスト"]
    sheet["B7"] = "権限まわりの設計が決まっておらず、実装に着手できていません"
    sheet["C7"] = "システム要件定義"
    sheet["D7"] = "高"
    sheet["E7"] = (TODAY + timedelta(days=7)).strftime("%Y/%m/%d")
    sheet["B8"] = "先週の作業は予定通り完了しました。ありがとうございました"
    sheet["B9"] = "帳票のレイアウトが決まらない。あと検証環境でエラーが出ています"
    buffer = io.BytesIO()
    book.save(buffer)

    analyzed = client.post(
        "/api/imports/analyze",
        files={"file": ("filled.xlsx", buffer.getvalue(), "application/octet-stream")},
    ).json()
    rows = client.post(
        "/api/imports/issue-rows",
        json={"token": analyzed["token"], "mapping": analyzed["mapping"]},
    ).json()["rows"]

    # 記入例1行 + 記入3行
    assert len(rows) == 4
    filled = next(r for r in rows if "権限まわり" in r["text"])
    assert filled["task_hint"] == "システム要件定義"  # 「関連するタスク・工程」列から
    assert filled["severity_hint"] == "高"
    assert filled["due_date"] == str(TODAY + timedelta(days=7))

    body = client.post(
        "/api/issues/analyze", json={"project_id": project.id, "rows": rows}
    ).json()
    items = {item["statement"]: item for item in body["items"]}

    auth = next(item for key, item in items.items() if "権限まわり" in key)
    assert auth["label"] == "issue"
    assert auth["severity"] == "high"  # 記入された重要度が使われる
    assert any("記入された重要度" in reason for reason in auth["reasons"])
    assert auth["due_date"] == str(TODAY + timedelta(days=7))
    assert [c["title"] for c in auth["related_task_candidates"]] == ["システム要件定義"]

    assert any(item["label"] == "not_issue" for item in body["items"])  # お礼・報告は除外
    assert len([i for i in body["items"] if "帳票" in i["statement"] or "検証環境" in i["statement"]]) == 2  # 分割

    created = client.post(
        "/api/issues/bulk_create",
        json={
            "project_id": project.id,
            "items": [
                {
                    "title": auth["title"],
                    "description": auth["description"],
                    "severity": auth["severity"],
                    "due_date": auth["due_date"],
                }
            ],
        },
    ).json()
    assert created["created"] == 1
    issue = client.get(f"/api/projects/{project.id}/issues").json()[0]
    assert issue["severity"] == "high"
    assert issue["due_date"] == str(TODAY + timedelta(days=7))


def test_template_download_endpoint(client):
    response = client.get("/api/imports/template")
    assert response.status_code == 200
    assert "issue_list_template.xlsx" in response.headers["content-disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["課題リスト", "記入のしかた"]
