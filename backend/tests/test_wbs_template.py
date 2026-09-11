"""クライアント記入用のWBSテンプレートが、そのまま取り込めること。"""
from __future__ import annotations

import io
from datetime import timedelta

from openpyxl import load_workbook

from app.importer.wbs_template import build_wbs_template, wbs_template_bytes

from .conftest import TODAY


def test_template_has_a_goal_column_and_examples():
    book = build_wbs_template(TODAY)
    sheet = book["WBS"]

    headers = [cell.value for cell in sheet[5]]
    assert headers[:6] == [
        "No", "大カテゴリ", "このカテゴリのゴール", "タスク（必須）", "タスクの詳細", "担当者",
    ]
    # 記入例は2行。同じカテゴリの2行目はゴールが空欄（1行書けばよいことを示す）
    assert sheet["B6"].value == "要件定義" and sheet["B7"].value == "要件定義"
    assert sheet["C6"].value and not sheet["C7"].value
    assert sheet["A6"].font.italic is True
    assert sheet["D8"].value is None  # 記入欄は空

    guide = book["記入のしかた"]
    text = "\n".join(str(row[0].value or "") for row in guide.iter_rows(max_col=1))
    assert "カテゴリごとに、達成したい状態を1行で書いてください" in text
    assert "必須は「タスク（必須）」の列だけです" in text


def test_template_offers_a_status_dropdown_and_no_progress_column():
    book = build_wbs_template(TODAY)
    sheet = book["WBS"]
    validations = {dv.type: dv for dv in sheet.data_validations.dataValidation}
    assert validations["list"].formula1 == '"未着手,進行中,完了,保留"'

    # 進捗率(%)は書いてもらわない。状態だけで足りる
    headers = [cell.value for cell in sheet[5]]
    assert "進捗率" not in headers
    assert "whole" not in validations
    guide = "\n".join(
        str(row[0].value or "") for row in book["記入のしかた"].iter_rows(max_col=1)
    )
    assert "％での進捗率は書かなくて構いません" in guide


def test_template_is_recognised_as_a_wbs(client):
    analyzed = client.post(
        "/api/imports/analyze",
        files={"file": ("wbs.xlsx", wbs_template_bytes(TODAY), "application/octet-stream")},
    ).json()

    assert analyzed["header_row"] == 4  # タイトル3行の下
    assert analyzed["content_kind"] == "wbs"
    mapping = analyzed["mapping"]
    assert mapping["title"] == "タスク(必須)"  # 全角括弧はNFKCで半角化
    assert mapping["parent"] == "大カテゴリ"
    assert mapping["category_goal"] == "このカテゴリのゴール"
    assert mapping["description"] == "タスクの詳細"
    assert mapping["owner"] == "担当者"
    assert mapping["planned_end"] == "期限"
    # 課題列は持たせていない（課題は課題リストのテンプレートで集める）
    assert mapping["issue"] is None
    assert analyzed["warnings"] == []


def test_a_filled_in_template_keeps_the_category_goal(client):
    book = build_wbs_template(TODAY)
    sheet = book["WBS"]
    rows = [
        ("1", "設計", "画面と帳票の仕様が確定し、実装に着手できる状態", "基本設計書作成", "佐藤", 7, "進行中"),
        ("2", "設計", "", "設計レビュー", "佐藤", 14, "未着手"),
        ("3", "開発", "本番相当のデータで主要業務が一通り動く状態", "共通機能実装", "田中", 30, ""),
        ("4", "開発", "", "環境構築", "田中", -3, "完了"),
    ]
    for offset, (code, category, goal, title, owner, due, status) in enumerate(rows):
        row = 8 + offset
        sheet.cell(row=row, column=1, value=code)
        sheet.cell(row=row, column=2, value=category)
        sheet.cell(row=row, column=3, value=goal)
        sheet.cell(row=row, column=4, value=title)
        sheet.cell(row=row, column=6, value=owner)
        sheet.cell(row=row, column=8, value=(TODAY + timedelta(days=due)).strftime("%Y/%m/%d"))
        sheet.cell(row=row, column=9, value=status)
    buffer = io.BytesIO()
    book.save(buffer)

    analyzed = client.post(
        "/api/imports/analyze",
        files={"file": ("filled.xlsx", buffer.getvalue(), "application/octet-stream")},
    ).json()
    committed = client.post(
        "/api/imports/commit",
        json={
            "token": analyzed["token"],
            "sheet": analyzed["selected_sheet"],
            "header_row": analyzed["header_row"],
            "mapping": analyzed["mapping"],
            "new_project_name": "WBSテンプレート",
            "match_by": "none",
        },
    ).json()

    tasks = client.get(f"/api/projects/{committed['project_id']}/tasks").json()
    by_title = {task["title"]: task for task in tasks}

    # カテゴリ（親タスク）が作られ、ゴールがその説明として残る
    design = by_title["設計"]
    assert design["is_summary"] is True
    assert design["description"] == "画面と帳票の仕様が確定し、実装に着手できる状態"
    # ゴールは同じカテゴリの1行目だけで足りる
    assert by_title["設計レビュー"]["parent_task_id"] == design["id"]
    assert by_title["開発"]["description"] == "本番相当のデータで主要業務が一通り動く状態"
    # タスク側の詳細はゴールで上書きされない
    assert by_title["基本設計書作成"]["description"] is None
    assert by_title["基本設計書作成"]["owner_name"] == "佐藤"
    assert by_title["基本設計書作成"]["planned_end"] == str(TODAY + timedelta(days=7))
    # 進捗率の列は無い。「完了」だけは100%と言い切れるので入れる
    assert by_title["環境構築"]["progress"] == 100.0
    assert by_title["基本設計書作成"]["progress"] == 0.0


def test_second_import_updates_the_goal(client):
    """ゴールを書き直した同じファイルを投げ直せば、カテゴリの説明も更新される。"""
    first = build_wbs_template(TODAY)["WBS"]

    def upload(goal: str) -> dict:
        book = build_wbs_template(TODAY)
        sheet = book["WBS"]
        sheet.cell(row=8, column=1, value="1")
        sheet.cell(row=8, column=2, value="テスト")
        sheet.cell(row=8, column=3, value=goal)
        sheet.cell(row=8, column=4, value="単体テスト")
        buffer = io.BytesIO()
        book.save(buffer)
        return client.post(
            "/api/imports/analyze",
            files={"file": ("wbs.xlsx", buffer.getvalue(), "application/octet-stream")},
        ).json()

    assert first["C8"].value is None  # テンプレート自体は空欄

    analyzed = upload("現場が操作して本番移行を判断できる状態")
    committed = client.post(
        "/api/imports/commit",
        json={
            "token": analyzed["token"], "sheet": analyzed["selected_sheet"],
            "header_row": analyzed["header_row"], "mapping": analyzed["mapping"],
            "new_project_name": "WBS更新", "match_by": "none",
        },
    ).json()
    project_id = committed["project_id"]

    analyzed = upload("重大バグが0件で、現場が本番移行を判断できる状態")
    client.post(
        "/api/imports/commit",
        json={
            "token": analyzed["token"], "sheet": analyzed["selected_sheet"],
            "header_row": analyzed["header_row"], "mapping": analyzed["mapping"],
            "project_id": project_id, "match_by": "auto",
        },
    )

    tasks = client.get(f"/api/projects/{project_id}/tasks").json()
    category = next(task for task in tasks if task["title"] == "テスト")
    assert category["description"] == "重大バグが0件で、現場が本番移行を判断できる状態"


def test_template_download_endpoint(client):
    response = client.get("/api/imports/template?kind=wbs")
    assert response.status_code == 200
    assert "wbs_template.xlsx" in response.headers["content-disposition"]
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["WBS", "記入のしかた"]

    # 既定は課題リスト（既存の導線を壊さない）
    assert "issue_list_template.xlsx" in client.get(
        "/api/imports/template"
    ).headers["content-disposition"]
    assert client.get("/api/imports/template?kind=unknown").status_code == 422
