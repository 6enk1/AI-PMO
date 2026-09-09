"""クライアント記入用の課題リストテンプレート。

「体裁を整えて出してください」と頼んでも整ってこないのが前提なので、
テンプレート側は**必須列を1つだけ**にして、書きやすさを優先している。
残りの列は空欄でよく、埋まっていれば取り込み時にそのまま使う。
"""
from __future__ import annotations

import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Meiryo"
SHEET_NAME = "課題リスト"
GUIDE_SHEET = "記入のしかた"

HEADERS = [
    ("No", 6),
    ("課題・気になっていること（必須）", 62),
    ("関連するタスク・工程", 24),
    ("重要度", 10),
    ("期限", 13),
    ("記入者", 14),
    ("備考", 28),
]
INPUT_ROWS = 40

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="EAF1FB")   # 記入するセル
EXAMPLE_FILL = PatternFill("solid", fgColor="F2F2F2")  # 記入例
THIN = Side(style="thin", color="B4C6E7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

EXAMPLE_ROW = [
    "例",
    "外部連携APIの認証方式が決まっておらず、実装に着手できていません。先方の回答待ちです。",
    "システム要件定義",
    "高",
    None,  # 日付は実行時に入れる
    "山田",
    "10/3の定例で再度確認予定",
]

GUIDE_LINES = [
    ("この表について", True),
    ("・気になっていること、困っていること、判断してほしいことを、思いつくまま書いてください。", False),
    ("・体裁は整えなくて構いません。文章のままで大丈夫です。", False),
    ("", False),
    ("記入するところ", True),
    ("・水色のセルに入力してください。グレーの行は記入例です（消しても、残したままでも構いません）。", False),
    ("・必須は「課題・気になっていること」の列だけです。ほかは空欄で構いません。", False),
    ("・1行に1つの課題を書くのが理想ですが、1つのセルに複数書いても取り込み時に自動で分けます。", False),
    ("・「順調です」「完了しました」のような報告が混ざっていても構いません。取り込み時に自動で除外します。", False),
    ("", False),
    ("各列の説明", True),
    ("・No … 通し番号。空欄で構いません。", False),
    ("・課題・気になっていること（必須） … 起きていること、困っていることを文章で。", False),
    ("・関連するタスク・工程 … 分かれば。「要件定義」「テスト」など、工程名でも構いません。", False),
    ("・重要度 … 高 / 中 / 低 から選択。分からなければ空欄のままで構いません。", False),
    ("・期限 … いつまでに判断・対応が必要か。分からなければ空欄で構いません。", False),
    ("・記入者 … 誰が書いたか。空欄で構いません。", False),
    ("・備考 … 補足があれば。", False),
    ("", False),
    ("列や行を増やしても大丈夫です", True),
    ("・行は足りなければ追加してください。列を増やしても取り込めます。", False),
    ("・列の順番を入れ替えても構いません。取り込み時に列名から自動で判別します。", False),
    ("", False),
    ("取り込んだあとの流れ", True),
    ("・AI PMO の Excel Import から、このファイルをそのままアップロードしてください。", False),
    ("・1行ずつ「課題 / 要確認 / 課題ではない」に自動で仕分けされ、確認画面が表示されます。", False),
    ("・内容を確認・修正し、チェックを入れた行だけが課題として登録されます。", False),
]


def build_template(today: date | None = None) -> Workbook:
    today = today or date.today()
    book = Workbook()

    sheet = book.active
    sheet.title = SHEET_NAME

    sheet["A1"] = "課題リスト"
    sheet["A1"].font = Font(name=FONT, size=14, bold=True)
    sheet["A2"] = "気になっていること・困っていることを、文章のまま書いてください。体裁を整える必要はありません。"
    sheet["A2"].font = Font(name=FONT, size=9, color="595959")
    sheet["A3"] = "記入するのは水色のセルです。必須は「課題・気になっていること」の列だけで、ほかは空欄で構いません。"
    sheet["A3"].font = Font(name=FONT, size=9, color="595959")

    header_row = 5
    for index, (title, width) in enumerate(HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
        cell = sheet.cell(row=header_row, column=index, value=title)
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    sheet.row_dimensions[header_row].height = 28

    # 記入例（グレー・斜体）。実データと見分けが付くようにしておく。
    example = list(EXAMPLE_ROW)
    example[4] = (today + timedelta(days=14)).strftime("%Y/%m/%d")
    for index, value in enumerate(example, start=1):
        cell = sheet.cell(row=header_row + 1, column=index, value=value)
        cell.font = Font(name=FONT, size=10, italic=True, color="7F7F7F")
        cell.fill = EXAMPLE_FILL
        cell.border = BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=index == 2)
    sheet.row_dimensions[header_row + 1].height = 34

    first_input = header_row + 2
    last_input = first_input + INPUT_ROWS - 1
    for row in range(first_input, last_input + 1):
        sheet.row_dimensions[row].height = 30
        for index in range(1, len(HEADERS) + 1):
            cell = sheet.cell(row=row, column=index)
            cell.font = Font(name=FONT, size=10)
            cell.fill = INPUT_FILL
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=index == 2)

    severity = DataValidation(
        type="list", formula1='"高,中,低"', allow_blank=True, showDropDown=False,
        promptTitle="重要度", prompt="高 / 中 / 低 から選んでください。分からなければ空欄で構いません。",
    )
    sheet.add_data_validation(severity)
    severity.add(f"D{first_input}:D{last_input}")

    due = DataValidation(
        type="date", operator="greaterThan", formula1="DATE(2000,1,1)", allow_blank=True,
        promptTitle="期限", prompt="いつまでに判断・対応が必要かを入れてください（例: 2026/10/31）。",
        errorTitle="日付を入れてください", error="2026/10/31 のような日付で入力してください。",
    )
    sheet.add_data_validation(due)
    due.add(f"E{first_input}:E{last_input}")

    issue_prompt = DataValidation(
        type="textLength", operator="greaterThan", formula1="0", allow_blank=True,
        promptTitle="課題・気になっていること",
        prompt="文章のままで構いません。1つのセルに複数書いても、取り込み時に自動で分けます。",
    )
    sheet.add_data_validation(issue_prompt)
    issue_prompt.add(f"B{first_input}:B{last_input}")

    sheet.freeze_panes = sheet.cell(row=first_input, column=1)

    guide = book.create_sheet(GUIDE_SHEET)
    guide.column_dimensions["A"].width = 96
    for index, (line, is_heading) in enumerate(GUIDE_LINES, start=1):
        cell = guide.cell(row=index, column=1, value=line)
        cell.font = Font(name=FONT, size=11 if is_heading else 10, bold=is_heading)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    return book


def template_bytes(today: date | None = None) -> bytes:
    buffer = io.BytesIO()
    build_template(today).save(buffer)
    return buffer.getvalue()
