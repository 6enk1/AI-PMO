"""テスト用WBSファイルを生成する。

日付は実行日を基準にした相対日付で書き出すので、いつ実行しても
「期限超過」「期限間近」といった状況がそのまま再現できる。

    python samples/generate_wbs_samples.py
"""
from __future__ import annotations

import csv
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_DIR = Path(__file__).resolve().parent
TODAY = date.today()

FONT = "Meiryo"  # 日本語WBSなので Arial ではなく Meiryo を使う
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
PHASE_FILL = PatternFill("solid", fgColor="DCE6F1")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADERS = [
    "No", "作業内容", "担当", "着手日", "期限", "実績開始日", "実績終了日", "工数(人日)",
    "進捗", "状態", "重要度", "先行作業", "課題", "備考", "MS",
]
WIDTHS = [8, 26, 12, 12, 12, 12, 12, 11, 8, 10, 9, 20, 26, 24, 5]

# 完了タスクの実績終了日を予定からずらす（予実差のサンプルとして）
ACTUAL_END_SHIFT: dict[str, int] = {"1.2": 1, "1.3": 3}


def d(offset: int | None) -> str:
    return "" if offset is None else (TODAY + timedelta(days=offset)).strftime("%Y/%m/%d")


# No, 作業内容, 担当, 着手, 期限, 工数, 進捗, 状態, 重要度, 先行作業, 課題, 備考, MS
V1_ROWS: list[list] = [
    ["1", "要件フェーズ", "", None, None, "", "", "", "", "", "", "", ""],
    ["1.1", "現行業務調査", "山田 太郎", -45, -35, 8, "100%", "完了", "高", "", "", "", ""],
    ["1.2", "業務要件ヒアリング", "山田 太郎", -34, -25, 10, "100%", "完了", "高", "現行業務調査", "", "", ""],
    ["1.3", "要件定義書作成", "佐々木 花子", -24, -12, 12, "100%", "完了", "最優先", "業務要件ヒアリング", "", "", ""],
    ["2", "設計フェーズ", "", None, None, "", "", "", "", "", "", "", ""],
    ["2.1", "基本設計", "佐々木 花子", -12, -2, 10, "95%", "進行中", "高", "要件定義書作成", "", "レビュー指摘の反映待ち", ""],
    ["2.2", "API仕様確定", "佐々木 花子", -10, -4, 6, "60%", "進行中", "最優先", "要件定義書作成", "顧客からのAPI仕様回答待ち", "週次定例で再依頼済み", ""],
    ["2.3", "DB設計", "鈴木 大輔", -10, -1, 8, "80%", "進行中", "高", "基本設計", "", "", ""],
    ["2.4", "画面設計", "", -8, 2, 9, "40%", "進行中", "中", "基本設計", "", "担当者が未定のまま進行", ""],
    ["3", "開発フェーズ", "", None, None, "", "", "", "", "", "", "", ""],
    ["3.1", "共通基盤構築", "田中 美咲", -5, 5, 7, "30%", "進行中", "高", "DB設計", "", "", ""],
    ["3.2", "バックエンド実装", "田中 美咲", 0, 18, 20, "0%", "未着手", "高", "API仕様確定", "", "", ""],
    ["3.3", "フロントエンド実装", "田中 美咲", 2, 20, 18, "0%", "未着手", "高", "画面設計", "", "", ""],
    ["3.4", "バッチ開発", "田中 美咲", 5, 22, 12, "0%", "未着手", "中", "DB設計", "", "", ""],
    ["3.5", "外部連携開発", "田中 美咲", 8, 25, 15, "0%", "未着手", "最優先", "API仕様確定", "接続先の仕様書が未入手", "", ""],
    ["4", "テストフェーズ", "", None, None, "", "", "", "", "", "", "", ""],
    ["4.1", "単体テスト", "高橋 由紀", 20, 28, 10, "0%", "未着手", "中", "バックエンド実装", "", "", ""],
    ["4.2", "結合テスト", "高橋 由紀", 28, 38, 14, "0%", "未着手", "高", "単体テスト", "", "", ""],
    ["4.3", "総合テスト", "高橋 由紀", 38, 50, 16, "0%", "未着手", "高", "結合テスト", "", "", "○"],
    ["4.4", "受入テスト支援", "", 50, 58, 8, "0%", "未着手", "中", "総合テスト", "", "顧客側の体制が未確定", ""],
    ["5", "移行・リリース", "", None, None, "", "", "", "", "", "", "", ""],
    ["5.1", "データ移行設計", "鈴木 大輔", -3, 8, 9, "10%", "進行中", "高", "DB設計", "移行元データに重複レコードあり", "", ""],
    ["5.2", "移行リハーサル", "鈴木 大輔", 40, 48, 6, "0%", "未着手", "高", "データ移行設計", "", "", ""],
    ["5.3", "本番リリース", "山田 太郎", 58, 60, 3, "0%", "未着手", "最優先", "受入テスト支援", "", "", "○"],
    ["6", "プロジェクト管理", "", None, None, "", "", "", "", "", "", "", ""],
    ["6.1", "週次進捗会議", "山田 太郎", -45, 60, 12, "50%", "進行中", "低", "", "", "", ""],
    ["6.2", "ベンダー選定", "", -6, 1, 4, "20%", "進行中", "中", "", "見積が1社しか届いていない", "", ""],
]

# 1週間後の更新版。同期インポート（更新4件 / 追加2件）を試すためのファイル。
V2_OVERRIDES: dict[str, dict[int, object]] = {
    "2.1": {6: "100%", 7: "完了"},
    "2.2": {4: 3, 6: "85%"},           # 期限を再設定し、進捗も更新
    "2.3": {6: "100%", 7: "完了"},
    "2.4": {2: "佐々木 花子", 6: "55%"},  # 担当が決まった
}
V2_ADDED: list[list] = [
    ["3.6", "性能テスト環境構築", "田中 美咲", 10, 24, 6, "0%", "未着手", "中", "共通基盤構築", "", "今回追加", ""],
    ["4.5", "移行データ検証", "鈴木 大輔", 45, 52, 5, "0%", "未着手", "中", "移行リハーサル", "", "今回追加", ""],
]

EN_ROWS = [
    ["1", "Requirements Analysis", "Alice Chen", -30, -20, "100%", "Done", "High", "", ""],
    ["2", "Solution Design", "Alice Chen", -19, -5, "70%", "In Progress", "High", "Requirements Analysis", "Design review pending"],
    ["3", "Backend Development", "Bob Smith", -4, 14, "20%", "In Progress", "High", "Solution Design", ""],
    ["4", "Frontend Development", "Carol Diaz", 0, 18, "0%", "Not Started", "Medium", "Solution Design", ""],
    ["5", "Integration Testing", "", 19, 30, "0%", "Not Started", "High", "Backend Development", "QA resource not assigned"],
    ["6", "Go Live", "Bob Smith", 31, 33, "0%", "Not Started", "Critical", "Integration Testing", ""],
]


def _style_sheet(sheet, title_rows: int) -> None:
    header_row = title_rows + 1
    for index, width in enumerate(WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)

    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=len(HEADERS)):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=False)

    for cell in sheet[header_row]:
        cell.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER

    for row in sheet.iter_rows(min_row=header_row + 1, max_row=sheet.max_row, max_col=len(HEADERS)):
        is_phase = "." not in str(row[0].value or "")
        for cell in row:
            cell.border = BORDER
            if is_phase:
                cell.fill = PHASE_FILL
                cell.font = Font(name=FONT, size=10, bold=True)


def _actual_dates(row: list) -> tuple[str, str]:
    """状態から実績開始日・実績終了日を組み立てる（未着手は空欄）。"""
    status = str(row[7] or "")
    start, end = row[3], row[4]
    if status == "完了":
        shift = ACTUAL_END_SHIFT.get(str(row[0]), 0)
        return d(start), d(end + shift if isinstance(end, int) else None)
    if status == "進行中":
        return d(start), ""
    return "", ""


def _write_sheet(sheet, rows: list[list], subtitle: str) -> None:
    sheet["A1"] = "基幹システム刷新プロジェクト WBS"
    sheet["A1"].font = Font(name=FONT, size=13, bold=True)
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(name=FONT, size=9, color="808080")
    sheet.append([])  # 3行目は空行（ヘッダー自動検出のテストを兼ねる）
    sheet.append(HEADERS)
    for row in rows:
        actual_start, actual_end = _actual_dates(row)
        sheet.append(
            [
                row[0], row[1], row[2],
                d(row[3]) if isinstance(row[3], int) else "",
                d(row[4]) if isinstance(row[4], int) else "",
                actual_start, actual_end,
                row[5], row[6], row[7], row[8], row[9], row[10], row[11], row[12],
            ]
        )
    _style_sheet(sheet, title_rows=3)


def build_v1() -> Path:
    book = Workbook()
    sheet = book.active
    sheet.title = "WBS"
    _write_sheet(sheet, V1_ROWS, f"作成日: {TODAY:%Y/%m/%d}　※AI PMO 動作確認用のサンプル（初版）")
    path = OUT_DIR / "wbs_sample_v1.xlsx"
    book.save(path)
    return path


def build_v2() -> Path:
    rows: list[list] = []
    for row in V1_ROWS:
        updated = list(row)
        for column, value in V2_OVERRIDES.get(str(row[0]), {}).items():
            updated[column] = value
        rows.append(updated)
    # 追加行はフェーズごとの末尾に差し込む
    for added in V2_ADDED:
        phase = str(added[0]).split(".")[0]
        last = max(i for i, r in enumerate(rows) if str(r[0]).split(".")[0] == phase)
        rows.insert(last + 1, added)

    book = Workbook()
    sheet = book.active
    sheet.title = "WBS"
    _write_sheet(sheet, rows, f"更新日: {TODAY:%Y/%m/%d}　※同期インポート確認用（更新4件・追加2件）")
    path = OUT_DIR / "wbs_sample_v2_update.xlsx"
    book.save(path)
    return path


def build_csv() -> Path:
    path = OUT_DIR / "wbs_sample_en.csv"
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(
            ["ID", "Task Name", "Assignee", "Start Date", "Due Date", "Progress",
             "Status", "Priority", "Predecessor", "Issue"]
        )
        for row in EN_ROWS:
            writer.writerow([row[0], row[1], row[2], d(row[3]), d(row[4]), *row[5:]])
    return path


def main() -> None:
    for path in (build_v1(), build_v2(), build_csv()):
        print(f"created: {path}")


if __name__ == "__main__":
    main()
